from __future__ import annotations

import csv
import datetime as dt
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

MAX_ROWS_PER_PAGE = 40
MAX_COLUMNS_PER_ROW = 30
MAX_CELL_CHARS = 240


def _clean_text(value: Any, *, max_chars: int = MAX_CELL_CHARS) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        text = value.isoformat()
    else:
        text = str(value)
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return ""
    if len(cleaned) > max_chars:
        return cleaned[: max_chars - 3].rstrip() + "..."
    return cleaned


def _sanitize_sheet_name(sheet_name: str) -> str:
    sanitized = re.sub(r"\s+", "_", str(sheet_name).strip())
    sanitized = re.sub(r"[^a-zA-Z0-9_\-]", "", sanitized)
    return sanitized or "Sheet"


def _excel_col_label(column_index: int) -> str:
    index = max(1, int(column_index))
    chars: List[str] = []
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        chars.append(chr(ord("A") + remainder))
    return "".join(reversed(chars))


def _detect_row_content_type(row_cells: Iterable[Dict[str, Any]]) -> str:
    has_formula = False
    cell_count = 0
    for cell in row_cells:
        cell_count += 1
        if cell.get("is_formula"):
            has_formula = True
            break
    if has_formula:
        return "FORMULA"
    if cell_count >= 2:
        return "TABLE"
    return "TEXT"


def _build_row_text(row_cells: List[Dict[str, Any]]) -> str:
    parts = []
    for cell in row_cells:
        cell_ref = str(cell.get("cell_ref") or "")
        value = str(cell.get("value") or "")
        if not value:
            continue
        parts.append(f"{cell_ref}: {value}")
    return " | ".join(parts)


def _chunk_rows(
    *,
    sheet_name: str,
    row_units: List[Dict[str, Any]],
    page_number_start: int,
    max_rows_per_page: int,
) -> Dict[str, Any]:
    pages: List[Dict[str, Any]] = []
    evidence_units: List[Dict[str, Any]] = []
    page_metadata: Dict[int, Dict[str, Any]] = {}

    page_number = page_number_start
    for chunk_start in range(0, len(row_units), max_rows_per_page):
        chunk = row_units[chunk_start : chunk_start + max_rows_per_page]
        if not chunk:
            continue

        page_lines = []
        for unit in chunk:
            unit_copy = dict(unit)
            unit_copy["page"] = page_number
            evidence_units.append(unit_copy)
            page_lines.append(
                f"[EVID type={unit_copy['content_type']} element={unit_copy['element_id']}] {unit_copy['text']}"
            )

        row_numbers = [int(unit.get("row") or 0) for unit in chunk if int(unit.get("row") or 0) > 0]
        row_start = min(row_numbers) if row_numbers else 0
        row_end = max(row_numbers) if row_numbers else 0

        pages.append({"page": page_number, "text": "\n".join(page_lines)})
        page_metadata[page_number] = {
            "sheet": sheet_name,
            "row_start": row_start,
            "row_end": row_end,
            "row_count": len(chunk),
        }
        page_number += 1

    return {
        "pages": pages,
        "evidence_units": evidence_units,
        "page_metadata": page_metadata,
        "next_page_number": page_number,
    }


def _build_sheet_row_units_from_csv(csv_path: str, *, max_columns_per_row: int) -> Dict[str, Any]:
    csv_file = Path(csv_path)
    sheet_name = _sanitize_sheet_name(csv_file.stem or "Sheet1")

    rows: List[List[str]] = []
    last_error: Optional[Exception] = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(csv_path, "r", encoding=encoding, newline="") as handle:
                rows = list(csv.reader(handle))
            last_error = None
            break
        except UnicodeDecodeError as exc:
            last_error = exc

    if last_error is not None:
        raise ValueError(f"Unable to decode CSV file: {last_error}") from last_error

    row_units: List[Dict[str, Any]] = []
    for row_index, row in enumerate(rows, start=1):
        row_cells: List[Dict[str, Any]] = []
        for column_index, value in enumerate(row[:max_columns_per_row], start=1):
            cleaned = _clean_text(value)
            if not cleaned:
                continue
            cell_ref = f"{_excel_col_label(column_index)}{row_index}"
            row_cells.append(
                {
                    "row": row_index,
                    "col": column_index,
                    "cell_ref": cell_ref,
                    "value": cleaned,
                    "is_formula": isinstance(value, str) and cleaned.startswith("="),
                }
            )

        if not row_cells:
            continue

        row_text = _build_row_text(row_cells)
        if not row_text:
            continue

        content_type = _detect_row_content_type(row_cells)
        element_id = f"sheet:{sheet_name}:row:{row_index}"
        row_units.append(
            {
                "unit_id": f"sheet:{sheet_name}:row:{row_index}",
                "sheet": sheet_name,
                "row": row_index,
                "content_type": content_type,
                "element_id": element_id,
                "text": row_text,
            }
        )

    return {
        "sheet_name": sheet_name,
        "row_units": row_units,
        "row_count": len(row_units),
    }


def _build_sheet_row_units_from_xlsx(
    xlsx_path: str,
    *,
    max_columns_per_row: int,
) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except ModuleNotFoundError as exc:  # pragma: no cover - exercised only when dependency missing
        raise RuntimeError(
            "openpyxl is required for XLSX parsing. Install openpyxl in the runtime environment."
        ) from exc

    workbook = load_workbook(filename=xlsx_path, read_only=True, data_only=False)
    sheet_results: List[Dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            sheet_name = _sanitize_sheet_name(worksheet.title)
            row_units: List[Dict[str, Any]] = []

            for row_index, row in enumerate(
                worksheet.iter_rows(min_row=1, max_col=max_columns_per_row),
                start=1,
            ):
                row_cells: List[Dict[str, Any]] = []
                for column_index, cell in enumerate(row, start=1):
                    raw_value = getattr(cell, "value", None)
                    cleaned = _clean_text(raw_value)
                    if not cleaned:
                        continue
                    is_formula = bool(getattr(cell, "data_type", "") == "f") or (
                        isinstance(raw_value, str) and cleaned.startswith("=")
                    )
                    cell_ref = f"{_excel_col_label(column_index)}{row_index}"
                    row_cells.append(
                        {
                            "row": int(row_index),
                            "col": int(column_index),
                            "cell_ref": cell_ref,
                            "value": cleaned,
                            "is_formula": is_formula,
                        }
                    )

                if not row_cells:
                    continue

                row_text = _build_row_text(row_cells)
                if not row_text:
                    continue

                content_type = _detect_row_content_type(row_cells)
                element_id = f"sheet:{sheet_name}:row:{row_index}"
                row_units.append(
                    {
                        "unit_id": f"sheet:{sheet_name}:row:{row_index}",
                        "sheet": sheet_name,
                        "row": int(row_index),
                        "content_type": content_type,
                        "element_id": element_id,
                        "text": row_text,
                    }
                )

            sheet_results.append(
                {
                    "sheet_name": sheet_name,
                    "row_units": row_units,
                    "row_count": len(row_units),
                }
            )
    finally:
        workbook.close()

    return sheet_results


def load_excel_multimodal_payload(
    file_path: str,
    *,
    max_rows_per_page: int = MAX_ROWS_PER_PAGE,
    max_columns_per_row: int = MAX_COLUMNS_PER_ROW,
) -> Dict[str, Any]:
    source_path = Path(file_path)
    suffix = source_path.suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        raise ValueError("Unsupported spreadsheet input. Expected .xlsx or .csv file.")

    if suffix == ".csv":
        sheet_results = [_build_sheet_row_units_from_csv(file_path, max_columns_per_row=max_columns_per_row)]
    else:
        sheet_results = _build_sheet_row_units_from_xlsx(
            file_path,
            max_columns_per_row=max_columns_per_row,
        )

    pages: List[Dict[str, Any]] = []
    page_metadata: Dict[int, Dict[str, Any]] = {}
    evidence_units: List[Dict[str, Any]] = []
    page_number = 1
    sheet_names: List[str] = []

    for sheet_result in sheet_results:
        row_units = sheet_result.get("row_units", [])
        sheet_name = str(sheet_result.get("sheet_name") or "Sheet")
        if not row_units:
            continue

        chunked = _chunk_rows(
            sheet_name=sheet_name,
            row_units=row_units,
            page_number_start=page_number,
            max_rows_per_page=max(1, int(max_rows_per_page)),
        )
        pages.extend(chunked["pages"])
        evidence_units.extend(chunked["evidence_units"])
        page_metadata.update(chunked["page_metadata"])
        page_number = int(chunked["next_page_number"])
        sheet_names.append(sheet_name)

    if not pages:
        raise ValueError("No usable rows found in spreadsheet input.")

    return {
        "document": {
            "source": source_path.name,
            "input_type": suffix.lstrip("."),
            "sheet_names": sheet_names,
            "sheet_count": len(sheet_names),
        },
        "pages": pages,
        "page_metadata": page_metadata,
        "evidence_units": evidence_units,
    }
